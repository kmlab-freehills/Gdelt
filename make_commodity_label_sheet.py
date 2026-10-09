"""
コモディティ（需要シグナル）の人手ラベル付け用Excelシートを作る（DBは読み取りのみ）

狭いクエリ（analyze）で集めた記事のうち、現在の需要プロンプトで処理済みのものから、
LLMの判定（除外／★2／★3）ごとに無作為に選ぶ。★3は少ないため多めに取り、
精度は各層の母数（meta.json）で重み付けして推定する。

出力（--out 配下の日時フォルダ）:
  - commodity_labels.xlsx : ラベル付け用シート
  - meta.json             : 層ごとの母数と標本の記事ID

実行例:
  docker compose run --rm --no-deps -v "$PWD:/app" llm_processor sh -c \
    "pip install -q openpyxl && python make_commodity_label_sheet.py"
"""

import argparse
import json
import os
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import text

from database import SessionLocal
from llm_processor import PROMPT_VERSIONS
from make_label_sheet import translate_titles

COMMODITIES = ["copper", "gold", "uranium", "rare earth", "natural gas"]
COMMODITY_JA = {"copper": "銅", "gold": "金", "uranium": "ウラン", "rare earth": "レアアース", "natural gas": "天然ガス"}
TONE_JA = {"bullish": "強気", "bearish": "弱気", "neutral": "中立"}
CHECK_CHOICES = ["正しい", "誤り", "記事に記載なし"]

# 層の名前 → 条件（llm_analysis に対するSQL）
STRATA = {
    "除外": "(llm_analysis->>'excluded')::boolean",
    "★2": "NOT (llm_analysis->>'excluded')::boolean AND llm_analysis->>'rating' = '2'",
    "★3": "NOT (llm_analysis->>'excluded')::boolean AND llm_analysis->>'rating' = '3'",
}

BASE_WHERE = """
collection_mode = 'analyze' AND is_llm_processed AND target = ANY(:targets)
AND llm_analysis->>'prompt_version' = :version
"""

SAMPLE_SQL = """
SELECT id, target, url, title, publish_date, source_domain, left(coalesce(body, ''), 400) AS body_head,
       llm_analysis
FROM articles WHERE {base} AND {cond}
ORDER BY md5(id::text || :seed) LIMIT :limit
"""

# (列名, 幅, 入力欄か)
COLUMNS = [
    ("No.", 5, False),
    ("記事ID", 7, False),
    ("品目", 9, False),
    ("LLMの判定", 8, False),
    ("観測日", 11, False),
    ("媒体", 18, False),
    ("見出し（日本語訳）", 40, False),
    ("見出し（原文）", 40, False),
    ("URL", 12, False),
    ("本文の冒頭", 50, False),
    ("LLMの判定理由", 45, False),
    ("LLM：方向", 8, False),
    ("LLM：きっかけ", 30, False),
    ("LLM：仕組み", 30, False),
    ("LLM：影響", 30, False),
    ("LLM：予想・記録との比較（引用）", 30, False),
    ("①需要に関わる事実がある記事か", 12, True),
    ("②正しい方向", 10, True),
    ("③きっかけ", 12, True),
    ("④仕組み", 12, True),
    ("⑤影響", 12, True),
    ("⑥あなたが付ける★", 11, True),
    ("メモ", 30, True),
]

INSTRUCTIONS = [
    ("このシートの目的", "コモディティの需要シグナルについて、LLMの判断（関係するか・方向・因果・重要度）が正しいかを確かめます。黄色の列だけ記入してください（プルダウンで選べます）。"),
    ("進め方", "URLを開いて記事を読みます。リンク切れの場合は「見出し」と「本文の冒頭」で判断し、メモに「リンク切れ」と書いてください。先頭がLLMが除外した記事（1件1分程度）、その後が採用した記事（1件3分程度）です。"),
    ("①需要に関わる事実がある記事か", "「はい」= その品目の需要（または需給バランス）に影響する具体的な事実を報じている記事（新規受注、増産・減産、政策、輸出規制、消費統計など）。「いいえ」= 価格の動きだけの市況、テクニカル分析、投資の勧め、一般論、関係ない話題。"),
    ("②正しい方向", "①が「はい」のときだけ。強気=需要が増える（または供給が締まる）方向の材料／弱気=需要が減る（または供給がだぶつく）方向の材料／中立=どちらとも言えない・混在。"),
    ("③きっかけ・④仕組み・⑤影響", "①が「はい」のときだけ。LLMの書いた内容が記事の内容と合っていれば「正しい」、違えば「誤り」。LLMが空欄で、記事にも書かれていなければ「記事に記載なし」。LLMが空欄なのに記事に書かれている場合は「誤り」にしてメモに書いてください。"),
    ("⑥あなたが付ける★", "①が「はい」のときだけ。LLMの★は見ずに、記事だけで付けてください。★3=記事自身が予想・コンセンサス・過去の記録と比べて「予想以上」「過去最高」「○年ぶり」などと明示している、または突発的な供給・需要ショック（輸出禁止、工場閉鎖、新たな義務化など）をその規模とともに報じている／★2=需要に影響しうる新しい具体的事実（予想との比較はない）／★1=既知の傾向の再確認、具体的な新事実のない論評。"),
    ("保存", "記入後はこのExcelファイルのまま上書き保存してください。"),
]


def main() -> None:
    parser = argparse.ArgumentParser(description="コモディティの人手ラベル付け用シートを作る")
    parser.add_argument("--excluded", type=int, default=10, help="除外の層から選ぶ件数")
    parser.add_argument("--rating2", type=int, default=12, help="★2の層から選ぶ件数")
    parser.add_argument("--rating3", type=int, default=10, help="★3の層から選ぶ件数")
    parser.add_argument("--seed", default="commodity-2026-10-07", help="標本を決める文字列")
    parser.add_argument("--out", default="eval_results/commodity_labels", help="出力先ディレクトリの親")
    parser.add_argument("--no-translate", action="store_true", help="見出しの日本語訳を付けない")
    args = parser.parse_args()

    sizes = {"除外": args.excluded, "★2": args.rating2, "★3": args.rating3}
    params = {"targets": COMMODITIES, "version": PROMPT_VERSIONS["demand"], "seed": args.seed}
    session = SessionLocal()
    try:
        population = {
            name: session.execute(text(f"SELECT count(*) FROM articles WHERE {BASE_WHERE} AND {cond}"), params).scalar()
            for name, cond in STRATA.items()
        }
        population["全体"] = session.execute(text(f"SELECT count(*) FROM articles WHERE {BASE_WHERE}"), params).scalar()
        rows = []
        for name, cond in STRATA.items():
            sql = SAMPLE_SQL.format(base=BASE_WHERE, cond=cond)
            for r in session.execute(text(sql), {**params, "limit": sizes[name]}).mappings().all():
                rows.append({**dict(r), "stratum": name})
    finally:
        session.close()
    print(f"母数 {population}、標本 {len(rows)} 件", flush=True)

    out_dir = os.path.join(args.out, datetime.now().strftime("%Y%m%d_%H%M%S"))
    os.makedirs(out_dir, exist_ok=True)
    meta = {
        "prompt_version": PROMPT_VERSIONS["demand"],
        "seed": args.seed,
        "population": population,
        "sample": {name: [r["id"] for r in rows if r["stratum"] == name] for name in STRATA},
        "period": {"from": min(r["publish_date"] for r in rows).isoformat(),
                   "to": max(r["publish_date"] for r in rows).isoformat()} if rows else None,
    }
    with open(os.path.join(out_dir, "meta.json"), "w", encoding="utf-8") as f:
        json.dump(meta, f, ensure_ascii=False, indent=2)

    translations = {} if args.no_translate else translate_titles({r["id"]: r["title"] or "" for r in rows})
    write_sheet(rows, translations, os.path.join(out_dir, "commodity_labels.xlsx"))
    print(f"出力先: {out_dir}")


def write_sheet(rows: list[dict], translations: dict[int, str], out: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "ラベル付け"
    header_fill = PatternFill("solid", fgColor="DDE5EC")
    input_header_fill = PatternFill("solid", fgColor="F2C94C")
    input_fill = PatternFill("solid", fgColor="FFF6D5")
    wrap = Alignment(wrap_text=True, vertical="top")

    for col, (name, width, is_input) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.font = Font(bold=True)
        cell.fill = input_header_fill if is_input else header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[cell.column_letter].width = width

    for n, r in enumerate(rows, 1):
        a = r["llm_analysis"] or {}
        causal = a.get("causal") or {}
        evidence = a.get("evidence") or {}
        kept = r["stratum"] != "除外"
        values = [
            n, r["id"], COMMODITY_JA.get(r["target"], r["target"]), r["stratum"],
            r["publish_date"].strftime("%Y-%m-%d") if r["publish_date"] else "",
            r["source_domain"], translations.get(r["id"], ""), r["title"], r["url"],
            (r["body_head"] or "").replace("\n", " "), a.get("reason") or "",
            TONE_JA.get(a.get("tone"), a.get("tone") or "") if kept else "",
            (causal.get("trigger") or "") if kept else "",
            (causal.get("mechanism") or "") if kept else "",
            (causal.get("effect") or "") if kept else "",
            (evidence.get("surprise_evidence") or "") if kept else "",
        ]
        for col, value in enumerate(values, 1):
            ws.cell(row=n + 1, column=col, value=value).alignment = wrap
        url_cell = ws.cell(row=n + 1, column=9)
        if r["url"]:
            url_cell.hyperlink = r["url"]
            url_cell.value = "記事を開く"
            url_cell.font = Font(color="1F5FA8", underline="single")
        for col, (_, _, is_input) in enumerate(COLUMNS, 1):
            if is_input:
                ws.cell(row=n + 1, column=col).fill = input_fill

    last = len(rows) + 1

    def add_list(col_name: str, choices: list[str]) -> None:
        col = next(i for i, (name, _, _) in enumerate(COLUMNS, 1) if name == col_name)
        letter = ws.cell(row=1, column=col).column_letter
        dv = DataValidation(type="list", formula1='"' + ",".join(choices) + '"', allow_blank=True)
        dv.error = "一覧から選んでください"
        ws.add_data_validation(dv)
        dv.add(f"{letter}2:{letter}{last}")

    add_list("①需要に関わる事実がある記事か", ["はい", "いいえ"])
    add_list("②正しい方向", list(TONE_JA.values()))
    add_list("③きっかけ", CHECK_CHOICES)
    add_list("④仕組み", CHECK_CHOICES)
    add_list("⑤影響", CHECK_CHOICES)
    add_list("⑥あなたが付ける★", ["★1", "★2", "★3"])
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=len(COLUMNS)).column_letter}{last}"

    guide = wb.create_sheet("記入方法", 0)
    guide.column_dimensions["A"].width = 26
    guide.column_dimensions["B"].width = 110
    for i, (title, body) in enumerate(INSTRUCTIONS, 1):
        guide.cell(row=i, column=1, value=title).font = Font(bold=True)
        guide.cell(row=i, column=2, value=body).alignment = Alignment(wrap_text=True, vertical="top")
        guide.cell(row=i, column=1).alignment = Alignment(vertical="top")
    wb.active = 1  # 開いたときはラベル付けのシートを表示

    wb.save(out)
    print(f"{len(rows)} 件を出力しました: {out}")


if __name__ == "__main__":
    main()

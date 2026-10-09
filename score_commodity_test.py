"""
コモディティの検証用標本を採点する（DBは読み取りのみ）

正解のラベルは、人が付けたもの（A_独立 → B_確認 の順に優先）、なければ採点役（Claude）のものを使う。
評価対象は、標本を作った時のプロンプトの判定（raw.jsonl、v4）と、DBに保存済みの判定（v2）。
各記事は、標本を選んだ層（v4の判定）の母数で重み付けする。採点役と人の一致率（A_独立）も出す。

実行例:
  docker compose run --rm --no-deps -v "$PWD:/app" llm_processor sh -c \
    "pip install -q openpyxl && python score_commodity_test.py --dir eval_results/commodity_test/20261007_164440"
"""

import argparse
import json
import os

from openpyxl import load_workbook
from sqlalchemy import text

from database import SessionLocal
from eval_demand_prompt import predict
from judge_demand import kappa
from score_labels import wilson

TONE_JA = {"bullish": "強気", "bearish": "弱気", "neutral": "中立"}


def read_sheet(path: str, name: str) -> dict[int, dict]:
    ws = load_workbook(path, data_only=True)[name]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h) for h in rows[0]]
    return {int(float(d["記事ID"])): d for d in (dict(zip(header, r)) for r in rows[1:]) if d.get("記事ID")}


def main() -> None:
    parser = argparse.ArgumentParser(description="コモディティの検証用標本を採点する")
    parser.add_argument("--dir", required=True, help="make_commodity_test_set.py の出力フォルダ")
    parser.add_argument("--labels", default="check_sheet_filled.xlsx", help="記入済みの確認シート")
    args = parser.parse_args()

    meta = json.load(open(os.path.join(args.dir, "meta.json"), encoding="utf-8"))
    raw = {d["id"]: d["result"] for d in map(json.loads, open(os.path.join(args.dir, "raw.jsonl"), encoding="utf-8"))}
    judge = {d["id"]: d for d in map(json.loads, open(os.path.join(args.dir, "judge_claude.jsonl"), encoding="utf-8"))}
    sheet = os.path.join(args.dir, args.labels)
    part_a, part_b = read_sheet(sheet, "A_独立"), read_sheet(sheet, "B_確認")

    session = SessionLocal()
    try:
        arts = {r["id"]: dict(r) for r in session.execute(
            text("SELECT id, target, title, body, llm_analysis FROM articles WHERE id = ANY(:i)"),
            {"i": list(judge)}).mappings().all()}
    finally:
        session.close()

    stratum = {i: k for k, ids in meta["sample"].items() for i in ids}
    weight = {i: meta["population"][stratum[i]] / len(meta["sample"][stratum[i]]) for i in judge}

    # 正解: A_独立（①②⑥）> B_確認（①のみ。②⑥は採点役）> 採点役
    gold, source = {}, {}
    for i, j in judge.items():
        g = {"relevant": j["relevant"], "direction": j["direction"], "rating": j["rating"]}
        source[i] = "Claude"
        a, b = part_a.get(i, {}), part_b.get(i, {})
        if a.get("①需要に関わる事実がある記事か") in ("はい", "いいえ"):
            g = {"relevant": a["①需要に関わる事実がある記事か"], "direction": a.get("②正しい方向"),
                 "rating": a.get("⑥あなたが付ける★")}
            source[i] = "人(A)"
        elif b.get("①あなたの判断") in ("はい", "いいえ"):
            if b["①あなたの判断"] != j["relevant"]:
                g = {"relevant": b["①あなたの判断"], "direction": None, "rating": None}
            source[i] = "人(B)"
        gold[i] = g

    preds = {}
    for i in judge:
        p = predict(dict(raw[i]), arts[i])
        v2 = arts[i]["llm_analysis"] or {}
        preds[i] = {
            "v4": {"kept": not p["excluded"], "dir": TONE_JA.get(p["tone"]), "rating": f"★{p['rating']}" if p["rating"] else None},
            "v2": {"kept": not v2.get("excluded"), "dir": TONE_JA.get(v2.get("tone")),
                   "rating": f"★{v2['rating']}" if v2.get("rating") else None},
        }

    def score(key: str) -> dict:
        kept = [i for i in judge if preds[i][key]["kept"]]
        excluded = [i for i in judge if not preds[i][key]["kept"]]
        tp = [i for i in kept if gold[i]["relevant"] == "はい"]
        fn = [i for i in excluded if gold[i]["relevant"] == "はい"]
        w = lambda ids: sum(weight[i] for i in ids)  # noqa: E731
        dir_ids = [i for i in tp if gold[i]["direction"]]
        rat_ids = [i for i in tp if gold[i]["rating"]]
        return {
            "kept": len(kept),
            "precision": wilson(len(tp), len(kept)),
            "precision_weighted_pct": round(100 * w(tp) / w(kept), 1) if kept else None,
            "excluded_correct": wilson(len(excluded) - len(fn), len(excluded)),
            "recall_weighted_pct": round(100 * w(tp) / (w(tp) + w(fn)), 1) if tp or fn else None,
            "direction": wilson(sum(1 for i in dir_ids if preds[i][key]["dir"] == gold[i]["direction"]), len(dir_ids)),
            "rating": wilson(sum(1 for i in rat_ids if preds[i][key]["rating"] == gold[i]["rating"]), len(rat_ids)),
            "rating3_confirmed": f"{sum(1 for i in kept if preds[i][key]['rating'] == '★3' and gold[i]['rating'] == '★3')}"
                                 f"/{sum(1 for i in kept if preds[i][key]['rating'] == '★3')}",
        }

    # 採点役と人の一致（A_独立で人が①を付けた記事）
    a_ids = [i for i in part_a if i in judge and source[i] == "人(A)"]
    rel_pairs = [(judge[i]["relevant"], gold[i]["relevant"]) for i in a_ids]
    both = [i for i in a_ids if judge[i]["relevant"] == "はい" and gold[i]["relevant"] == "はい"]
    b_ids = [i for i in part_b if source.get(i) == "人(B)"]
    summary = {
        "n": len(judge),
        "label_source": {s: sum(1 for v in source.values() if v == s) for s in set(source.values())},
        "v4": score("v4"),
        "v2": score("v2"),
        "judge_vs_human_A": {
            "relevant": {**wilson(sum(1 for a, b in rel_pairs if a == b), len(rel_pairs)), "kappa": kappa(rel_pairs)},
            "disagree": [(i, judge[i]["relevant"], gold[i]["relevant"]) for i in a_ids if judge[i]["relevant"] != gold[i]["relevant"]],
            "direction": f"{sum(1 for i in both if judge[i]['direction'] == gold[i]['direction'])}/{len(both)}",
            "rating": f"{sum(1 for i in both if judge[i]['rating'] == gold[i]['rating'])}/{len(both)}",
        },
        "judge_vs_human_B": wilson(sum(1 for i in b_ids if judge[i]["relevant"] == gold[i]["relevant"]), len(b_ids)),
        "changed_by_human": [(i, judge[i]["relevant"], gold[i]["relevant"]) for i in judge if judge[i]["relevant"] != gold[i]["relevant"]],
    }
    with open(os.path.join(args.dir, "score.json"), "w", encoding="utf-8") as f:
        json.dump(summary, f, ensure_ascii=False, indent=2)
    print(json.dumps(summary, ensure_ascii=False, indent=1))


if __name__ == "__main__":
    main()

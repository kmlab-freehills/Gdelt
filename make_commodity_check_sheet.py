"""
採点役（Claude）の採点を人が確かめるためのExcelシートを作る（DBは読み取りのみ）

make_commodity_test_set.py の標本について、2つのシートを作る:
  A_独立 : 無作為に選んだ --independent 件。LLMの判定も採点役の判定も見せずに、人が①②⑥を付ける（採点役の一致率を測る）
  B_確認 : Aに入らなかった記事のうち、採点役と評価対象のLLMで①の判断が分かれた記事。採点役の判断と根拠を見せて、人が①を確かめる

実行例:
  docker compose run --rm --no-deps -v "$PWD:/app" llm_processor sh -c \
    "pip install -q openpyxl && python make_commodity_check_sheet.py --dir eval_results/commodity_test/20261007_164440"
"""

import argparse
import hashlib
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import text

from database import SessionLocal
from eval_demand_prompt import predict
from make_commodity_label_sheet import COMMODITY_JA, TONE_JA
from make_label_sheet import translate_titles

SQL = """
SELECT id, target, url, title, publish_date, source_domain, body, left(coalesce(body, ''), 400) AS body_head
FROM articles WHERE id = ANY(:ids)
"""

BASE_COLUMNS = [("No.", 5), ("記事ID", 7), ("品目", 9), ("観測日", 11), ("媒体", 18), ("見出し（日本語訳）", 40),
                ("見出し（原文）", 40), ("URL", 12), ("本文の冒頭", 50)]
A_INPUTS = [("①需要に関わる事実がある記事か", 12), ("②正しい方向", 10), ("⑥あなたが付ける★", 11), ("メモ", 30)]
B_INFO = [("Claudeの判断", 9), ("Claudeの根拠（引用）", 40), ("Claudeのメモ", 30), ("qwenの判断", 9)]
B_INPUTS = [("①あなたの判断", 12), ("メモ", 30)]

GUIDE = [
    ("このファイルの目的", "コモディティの判定を採点したClaude（採点役）が、人の判断とどれだけ一致するかを確かめます。2枚のシートがあります。"),
    ("A_独立（先にやる）", "無作為に選んだ記事です。サイドバーのGeminiは使わず、ご自分の判断だけで①②⑥を付けてください（採点役の一致率を独立に測るため）。①が「いいえ」なら②⑥は空欄で構いません。1件1〜3分程度。"),
    ("B_確認", "Claudeとqwenで「需要に関わる記事か」の判断が分かれた記事です。Claudeの判断と根拠を見て、①だけ記入してください。Geminiを使っても構いません。1件30秒〜1分程度。"),
    ("①の基準", "「はい」= その品目の需要（または需給バランス）に影響する具体的な事実を報じている（新規受注、増産・減産、政策、輸出規制、消費統計、供給の途絶など）。「いいえ」= 価格の動きだけの市況、テクニカル分析、投資の勧め、一般論、この品目と関係ない話題（別の品目の話も含む）。"),
    ("②の基準", "強気=需要が増える、または供給が締まる・途絶える材料／弱気=需要が減る、または供給がだぶつく材料／中立=どちらとも言えない・混在。"),
    ("⑥の基準", "★3=記事自身が予想・コンセンサス・過去の記録と比べて明示している、または突発的な供給・需要ショックをその規模とともに報じている／★2=需要に影響しうる新しい具体的事実／★1=既知の傾向の再確認、新事実のない論評。"),
    ("保存", "記入後はこのExcelファイルのまま上書き保存してください。"),
]


def _sheet(wb, title, columns, inputs, rows, values_of):
    ws = wb.create_sheet(title)
    header_fill, input_header_fill = PatternFill("solid", fgColor="DDE5EC"), PatternFill("solid", fgColor="F2C94C")
    input_fill, wrap = PatternFill("solid", fgColor="FFF6D5"), Alignment(wrap_text=True, vertical="top")
    all_cols = columns + inputs
    for col, (name, width) in enumerate(all_cols, 1):
        c = ws.cell(row=1, column=col, value=name)
        c.font, c.alignment = Font(bold=True), Alignment(wrap_text=True, vertical="center")
        c.fill = input_header_fill if col > len(columns) else header_fill
        ws.column_dimensions[c.column_letter].width = width
    for n, r in enumerate(rows, 1):
        for col, value in enumerate(values_of(n, r), 1):
            ws.cell(row=n + 1, column=col, value=value).alignment = wrap
        url_cell = ws.cell(row=n + 1, column=8)
        if r["url"]:
            url_cell.hyperlink, url_cell.value = r["url"], "記事を開く"
            url_cell.font = Font(color="1F5FA8", underline="single")
        for col in range(len(columns) + 1, len(all_cols) + 1):
            ws.cell(row=n + 1, column=col).fill = input_fill
    last = len(rows) + 1

    def add_list(name, choices):
        col = next(i for i, (c, _) in enumerate(all_cols, 1) if c == name)
        letter = ws.cell(row=1, column=col).column_letter
        dv = DataValidation(type="list", formula1='"' + ",".join(choices) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{letter}2:{letter}{last}")
    return ws, add_list


def main() -> None:
    parser = argparse.ArgumentParser(description="採点役の採点を人が確かめるシートを作る")
    parser.add_argument("--dir", required=True, help="make_commodity_test_set.py の出力フォルダ（judge_claude.jsonl を含む）")
    parser.add_argument("--independent", type=int, default=10, help="A_独立 に入れる件数")
    args = parser.parse_args()

    judge = {d["id"]: d for d in map(json.loads, open(os.path.join(args.dir, "judge_claude.jsonl"), encoding="utf-8"))}
    raw = {d["id"]: d["result"] for d in map(json.loads, open(os.path.join(args.dir, "raw.jsonl"), encoding="utf-8"))}
    session = SessionLocal()
    try:
        arts = {r["id"]: dict(r) for r in session.execute(text(SQL), {"ids": list(judge)}).mappings().all()}
    finally:
        session.close()

    order = sorted(judge, key=lambda i: hashlib.md5(f"independent{i}".encode()).hexdigest())
    part_a = order[:args.independent]
    qwen_kept = {i: not predict(dict(raw[i]), arts[i])["excluded"] for i in judge}
    part_b = [i for i in order[args.independent:] if qwen_kept[i] != (judge[i]["relevant"] == "はい")]
    translations = translate_titles({i: arts[i]["title"] or "" for i in part_a + part_b})

    def base(n, a):
        return [n, a["id"], COMMODITY_JA.get(a["target"], a["target"]), a["publish_date"].strftime("%Y-%m-%d"),
                a["source_domain"], translations.get(a["id"], ""), a["title"], a["url"],
                (a["body_head"] or "").replace("\n", " ") or "（本文なし：見出しで判断）"]

    wb = Workbook()
    guide = wb.active
    guide.title = "記入方法"
    guide.column_dimensions["A"].width, guide.column_dimensions["B"].width = 22, 110
    for i, (t, b) in enumerate(GUIDE, 1):
        guide.cell(row=i, column=1, value=t).font = Font(bold=True)
        guide.cell(row=i, column=2, value=b).alignment = Alignment(wrap_text=True, vertical="top")

    _, add_a = _sheet(wb, "A_独立", BASE_COLUMNS, A_INPUTS, [arts[i] for i in part_a],
                      lambda n, a: base(n, a))
    add_a("①需要に関わる事実がある記事か", ["はい", "いいえ"])
    add_a("②正しい方向", list(TONE_JA.values()))
    add_a("⑥あなたが付ける★", ["★1", "★2", "★3"])

    _, add_b = _sheet(wb, "B_確認", BASE_COLUMNS + B_INFO, B_INPUTS, [arts[i] for i in part_b],
                      lambda n, a: base(n, a) + [judge[a["id"]]["relevant"], judge[a["id"]]["quote"],
                                                 judge[a["id"]]["note"], "採用" if qwen_kept[a["id"]] else "除外"])
    add_b("①あなたの判断", ["はい", "いいえ"])

    out = os.path.join(args.dir, "check_sheet.xlsx")
    wb.save(out)
    with open(os.path.join(args.dir, "check_meta.json"), "w", encoding="utf-8") as f:
        json.dump({"part_a": part_a, "part_b": part_b}, f, ensure_ascii=False, indent=2)
    print(f"A_独立 {len(part_a)} 件、B_確認 {len(part_b)} 件 → {out}")


if __name__ == "__main__":
    main()

"""
人手ラベル付け用のExcelシートを作る（DBは読み取りのみ）

export_labels.py が出力したCSVと同じ記事（同じ標本）について、日本語で記入できる
Excelファイルを作る。入力欄はプルダウンで選択式にし、見出しの日本語訳をLLMで付ける。
openpyxl が必要なため、一時コンテナでインストールしてから実行する。

実行例:
  docker compose run --rm --no-deps -v "$PWD:/app" llm_processor_construction sh -c \
    "pip install -q openpyxl && python make_label_sheet.py --csv eval_results/labels/labels_20261006_173604.csv"
"""

import argparse
import csv
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import text

from database import SessionLocal
from llm_processor import _OllamaBackend, _parse_json

STATUS_JA = {
    "planned": "計画", "groundbreaking": "着工", "under_construction": "施工中",
    "topped_out": "上棟", "completed": "竣工", "halted": "中断",
    "resumed": "再開", "cancelled": "中止", "unknown": "不明",
}
SOURCE_TYPES_JA = ["一次報道", "転載", "論説", "まとめ", "PR", "その他"]
CHECK_CHOICES = ["正しい", "誤り", "記事に記載なし"]

SQL = """
SELECT id, url, title, publish_date, source_domain, left(coalesce(body, ''), 400) AS body_head,
       (llm_analysis->>'excluded')::boolean AS excluded,
       llm_analysis->>'reason' AS reason,
       event_date,
       llm_analysis->'construction'->>'status'           AS status,
       llm_analysis->'construction'->>'building_name'    AS building_name,
       llm_analysis->'construction'->'location'->>'city' AS city,
       llm_analysis->'construction'->>'project_evidence' AS evidence
FROM articles WHERE id = ANY(:ids)
"""

# (列名, 幅, 入力欄か)
COLUMNS = [
    ("No.", 5, False),
    ("記事ID", 7, False),
    ("LLMの判定", 9, False),
    ("観測日", 11, False),
    ("媒体", 18, False),
    ("見出し（日本語訳）", 40, False),
    ("見出し（原文）", 40, False),
    ("URL", 12, False),
    ("本文の冒頭", 50, False),
    ("LLMの判定理由", 45, False),
    ("LLM：状態", 9, False),
    ("LLM：案件名", 22, False),
    ("LLM：都市", 12, False),
    ("LLM：日付", 11, False),
    ("LLM：根拠の引用", 30, False),
    ("①建設案件の記事か", 12, True),
    ("②正しい状態", 11, True),
    ("③案件名", 12, True),
    ("正しい案件名（誤りの場合）", 22, True),
    ("④都市", 12, True),
    ("正しい都市（誤りの場合）", 16, True),
    ("⑤日付", 12, True),
    ("正しい日付（誤りの場合）", 14, True),
    ("⑥ソースの種類", 11, True),
    ("メモ", 30, True),
]

INSTRUCTIONS = [
    ("このシートの目的", "LLMの判断が正しいかを確かめるため、記事を読んで正解を記入します。黄色の列だけ記入してください（プルダウンで選べます）。"),
    ("進め方", "URLを開いて記事を読みます。リンク切れの場合は「見出し」と「本文の冒頭」で判断し、メモに「リンク切れ」と書いてください。1〜50行目はLLMが除外した記事（1件1分程度）、51〜100行目は採用した記事（1件3分程度）です。除外の方から始めると早く進みます。"),
    ("①建設案件の記事か", "「はい」= 特定できる大規模建設の案件（ビル・工場・駅・道路・発電所など）について報じている記事。「いいえ」= それ以外（市況・政策一般・事件・関係ない話題など）。迷ったら「この記事だけで、衛星で見に行く場所が分かるか」で判断してください。"),
    ("②正しい状態", "①が「はい」のときだけ。計画=発表段階で未着工／着工=起工式・着工／施工中／上棟=最上部まで到達／竣工=完成・開業／中断／再開／中止／不明=記事から判断できない。"),
    ("③案件名・④都市・⑤日付", "LLMの値が合っていれば「正しい」、違えば「誤り」を選び、右隣の列に正しい値を書きます。記事にそもそも書かれていなければ「記事に記載なし」。①が「いいえ」の記事は空欄で構いません。"),
    ("⑥ソースの種類", "一次報道=その媒体の記者が取材した記事（関係者の話、現地の様子、資料の確認など）／転載=通信社や提携先の記事をそのまま載せたもの（「(AP)」「Reuters」など）／論説=コラム・社説・意見記事／まとめ=他媒体の報道の寄せ集め／PR=企業や自治体の発表をほぼそのまま載せたもの／その他=物件紹介・広告記事・告知など。①が「いいえ」の記事も記入してください。"),
    ("保存", "記入後はこのExcelファイルのまま上書き保存してください。"),
]


def translate_titles(titles: dict[int, str]) -> dict[int, str]:
    backend = _OllamaBackend()
    result = {}
    for i, (article_id, title) in enumerate(titles.items(), 1):
        prompt = (
            "Translate the following news headline into natural Japanese. "
            "If it is already Japanese, return it unchanged. "
            'Return ONLY valid JSON: {"ja": "<Japanese headline>"}\n\n'
            f"Headline: {title}"
        )
        try:
            result[article_id] = (_parse_json(backend.call(prompt)) or {}).get("ja") or ""
        except Exception as e:
            print(f"翻訳失敗 id={article_id}: {e}", flush=True)
            result[article_id] = ""
        if i % 20 == 0:
            print(f"翻訳 {i}/{len(titles)}", flush=True)
    return result


def main() -> None:
    parser = argparse.ArgumentParser(description="人手ラベル付け用のExcelシートを作る")
    parser.add_argument("--csv", required=True, help="export_labels.py が出力したCSV（同じ標本を使う）")
    parser.add_argument("--out", help="出力するxlsxのパス（省略時はCSVと同じ場所・同じ名前）")
    parser.add_argument("--no-translate", action="store_true", help="見出しの日本語訳を付けない")
    args = parser.parse_args()

    with open(args.csv, encoding="utf-8-sig") as f:
        ids = [int(r["id"]) for r in csv.DictReader(f)]

    session = SessionLocal()
    try:
        rows = {r["id"]: dict(r) for r in session.execute(text(SQL), {"ids": ids}).mappings().all()}
    finally:
        session.close()

    # 除外（1件1分程度）を先に、採用を後に並べる。それぞれの中はCSVの順番を保つ
    ordered = [rows[i] for i in ids if i in rows and rows[i]["excluded"]] + \
              [rows[i] for i in ids if i in rows and not rows[i]["excluded"]]
    translations = {} if args.no_translate else translate_titles({r["id"]: r["title"] or "" for r in ordered})
    out = args.out or os.path.splitext(args.csv)[0] + ".xlsx"
    write_sheet(ordered, translations, out)


def write_sheet(ordered: list[dict], translations: dict[int, str], out: str) -> None:
    """ラベル付け用のExcelを書き出す。

    ordered の各要素に必要なキー: id, excluded, publish_date, source_domain, title, url, body_head,
    reason, status（英語のコード）, building_name, city, event_date, evidence
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "ラベル付け"
    header_fill = PatternFill("solid", fgColor="DDE5EC")
    input_header_fill = PatternFill("solid", fgColor="F2C94C")
    input_fill = PatternFill("solid", fgColor="FFF6D5")
    wrap = Alignment(wrap_text=True, vertical="top")

    for col, (name, width, is_input) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.font = Font(bold=True)
        cell.fill = input_header_fill if is_input else header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[cell.column_letter].width = width

    for n, r in enumerate(ordered, 1):
        values = [
            n, r["id"], "除外" if r["excluded"] else "採用",
            r["publish_date"].strftime("%Y-%m-%d") if r["publish_date"] else "",
            r["source_domain"], translations.get(r["id"], ""), r["title"], r["url"],
            (r["body_head"] or "").replace("\n", " "), r["reason"],
            STATUS_JA.get(r["status"], r["status"] or "") if not r["excluded"] else "",
            r["building_name"] or "", r["city"] or "",
            r["event_date"].isoformat() if r["event_date"] else "", r["evidence"] or "",
        ]
        for col, value in enumerate(values, 1):
            cell = ws.cell(row=n + 1, column=col, value=value)
            cell.alignment = wrap
        url_cell = ws.cell(row=n + 1, column=8)
        if r["url"]:
            url_cell.hyperlink = r["url"]
            url_cell.value = "記事を開く"
            url_cell.font = Font(color="1F5FA8", underline="single")
        for col, (_, _, is_input) in enumerate(COLUMNS, 1):
            if is_input:
                ws.cell(row=n + 1, column=col).fill = input_fill

    last = len(ordered) + 1

    def add_list(col_name: str, choices: list[str]) -> None:
        col = next(i for i, (name, _, _) in enumerate(COLUMNS, 1) if name == col_name)
        letter = ws.cell(row=1, column=col).column_letter
        dv = DataValidation(type="list", formula1='"' + ",".join(choices) + '"', allow_blank=True)
        dv.error = "一覧から選んでください"
        ws.add_data_validation(dv)
        dv.add(f"{letter}2:{letter}{last}")

    add_list("①建設案件の記事か", ["はい", "いいえ"])
    add_list("②正しい状態", list(STATUS_JA.values()))
    add_list("③案件名", CHECK_CHOICES)
    add_list("④都市", CHECK_CHOICES)
    add_list("⑤日付", CHECK_CHOICES)
    add_list("⑥ソースの種類", SOURCE_TYPES_JA)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=len(COLUMNS)).column_letter}{last}"

    guide = wb.create_sheet("記入方法", 0)
    guide.column_dimensions["A"].width = 22
    guide.column_dimensions["B"].width = 110
    for i, (title, body) in enumerate(INSTRUCTIONS, 1):
        guide.cell(row=i, column=1, value=title).font = Font(bold=True)
        guide.cell(row=i, column=2, value=body).alignment = Alignment(wrap_text=True, vertical="top")
        guide.cell(row=i, column=1).alignment = Alignment(vertical="top")
    wb.active = 1  # 開いたときはラベル付けのシートを表示

    wb.save(out)
    print(f"{len(ordered)} 件を出力しました: {out}（除外 {sum(1 for r in ordered if r['excluded'])} 件を先頭に配置）")


if __name__ == "__main__":
    main()
